# pip install python-pptx
import base64
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333333333333334)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string('ffffff')

# accent
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.65), Inches(0.1), Inches(1))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor.from_string('147c6d')
shape.line.fill.background()
shape.name = "Accent bar"
shape.rotation = 0

# title
shape = slide.shapes.add_textbox(Inches(1), Inches(0.7), Inches(10), Inches(1.2))
tf = shape.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.text = "Arrange visually.\nAutomate precisely."
tf.word_wrap = True
for paragraph in tf.paragraphs:
    paragraph.font.size = Pt(32)
    paragraph.font.name = "Arial"
    paragraph.font.color.rgb = RGBColor.from_string('243249')
shape.name = "Headline"
shape.rotation = 0

# body
shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(2))
tf = shape.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.text = "Use this sample to test your workflow.\nEvery position is stored in inches.\nZoom is only a screen transformation."
tf.word_wrap = True
for paragraph in tf.paragraphs:
    paragraph.font.size = Pt(20)
    paragraph.font.name = "Arial"
    paragraph.font.color.rgb = RGBColor.from_string('243249')
shape.name = "Description"
shape.rotation = 0

# image
shape = slide.shapes.add_picture(BytesIO(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAZAAAADwCAIAAAChXqV1AAAFbklEQVR4nO3csXEUWRRAUbFFSBsSBICzNvY6BAAhbRAkgEMEGFSpxA5II2mm593/z/Fhel7Xu/1blHjz9fu3O4CCv259AQDnEiwgQ7CADMECMgQLyBAsIEOwgAzBAjIEC8gQLCBDsIAMwQIyBAvIECwgQ7CADMECMgQLyBAsIEOwgAzBAjIEC8gQLCBDsIAMwQIyBAvIECwgQ7CADMECMgQLyBAsIEOwgAzBAjIEC8gQLCDj7WGf9Pe/Hw/7LOBg//3z8YBPccICMgQLyBAsIEOwgAzBAjIEC8gQLCBDsIAMwQIyBAvIECwgQ7CADMECMgQLyDjuv5dhSZ/ev3vuH/nw+cs1roQdCBbP8II8nfOXSBhnEiyecJFInf8R4sUjBIvfO6BTj3+ucnFKsPjFrTp1Srk4JVjc3U3q1Cnl4p5g7W5yqv7n56XK1s4Ea1+hVD0kWzsTrB1FU/WQbO1JsPayQKoekq3d+NWcjSxWq3urfi9OOWFtYfmVdtTahBPW+pav1b19vum2nLBWtuECO2qtzQlrWRvW6t7O331tgrUmG2sCS/JKuBqLes/r4XqcsJaiVqfMZCWCtQ6b+ScmswzBWoSdfJz5rEGwVmAbz2FKCxCsPHt4PrOqE6w2G/hcJpYmWGF272XMrUuwqmzda5helGAl2bfXM8MiwQIyBKvH0eBSTDJHsGLs2GWZZ4tgldiuazDVEMECMgQrw0Hgesy2QrAabNS1mXCCYAEZghXg4X8Mc55PsKazRUcy7eEEC8gQrNE88I9n5pMJFpAhWHN51N+KyY8lWECGYA3lIX9b5j+TYAEZgjWRx/sE7sJAggVkCBaQIVjjeBOZw72YRrCADMECMgRrFu8g07gjowgWkCFYQIZgDeLtYyb3ZQ7BAjIEC8gQLCBDsKbwg5LJ3J0hBAvIECwgQ7CADMECMgQLyBCsEfwj1Hzu0QSCBWQIFpAhWECGYAEZggVkCBaQIVhAhmABGYIFZAgWkCFYQIZgARmCBWQIFpAhWECGYAEZgjXCh89fbn0JPME9mkCwgAzBAjIEC8gQLCBDsIAMwZrCP0JN5u4MIVhAhmABGYIFZAjWIH5QMpP7ModgARmCBWQI1izePqZxR0YRLCBDsIAMwRrHO8gc7sU0ggVkCBaQIVgTeROZwF0YSLCADMEayuP9tsx/JsECMgRrLg/5WzH5sQQLyBCs0Tzqj2fmkwkWkCFY03ngH8m0hxOsAFt0DHOeT7CADMFq8PC/NhNOEKwMG3U9ZlshWECGYJU4CFyDqYYIVoztuizzbBGsHjt2KSaZI1hAhmAlORq8nhkWCVaVfXsN04sSrDBb9zLm1iVYbXbvuUwsTbDybOD5zKpOsFZgD89hSgsQrEXYxseZzxoEax128k9MZhmCtRSbecpMVvL21hfAhf3cz0/v3936Qm5PqtbjhLUmu2oCSxKsZe28sTt/97V5JVzZhq+HUrU2J6z17bPD+3zTbTlhbWH5o5ZUbcIJayOrbvWq34tTTlh7WeyoJVW7EawdLZAtqdqTYO0rmi2p2plg7S6ULalCsLi7e9CCgeXSKe4JFr+YUy6d4pRg8Xu3KpdO8QjB4gkPC3KleIkUZxIsnuG0LC9ImDzxYoLFq6gPR/KrOUCGYAEZggVkCBaQIVhAhmABGYIFZAgWkCFYQIZgARmCBWQIFpAhWECGYAEZb75+/3brawA4ixMWkCFYQIZgARmCBWQIFpAhWECGYAEZggVkCBaQIVhAhmABGYIFZAgWkCFYQIZgARmCBWQIFpAhWECGYAEZggVkCBaQIVhAhmABGYIFZAgWkCFYQIZgARmCBWQIFpAhWECGYAEZggVkCBaQ8QNAbClKDIcRFgAAAABJRU5ErkJggg==")), Inches(7), Inches(2), Inches(4.5), Inches(2.7))
shape.name = "Embedded demo image"
shape.rotation = 0

# footer
shape = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11), Inches(0.4))
tf = shape.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.text = "SLIDE COORDINATE STUDIO  /  DEMO"
tf.word_wrap = True
for paragraph in tf.paragraphs:
    paragraph.font.size = Pt(12)
    paragraph.font.name = "Arial"
    paragraph.font.color.rgb = RGBColor.from_string('243249')
shape.name = "Footer"
shape.rotation = 0

prs.save("layout.pptx")
